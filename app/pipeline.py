import os
from pathlib import Path
from typing import Any, Dict, List, Literal, Optional

from pydantic import BaseModel, Field

from app.gemini_client import generate_json
from docx import Document


class RequirementsMap(BaseModel):
    target_role: str
    must_have_keywords: List[str]
    preferred_keywords: List[str]
    leadership_themes: List[str]
    metric_categories: List[str]
    formatting_rules: List[str]


class EvidenceItem(BaseModel):
    id: str
    bullet: str
    metric_types: List[str]
    tags: List[str]
    source_files: List[str]


class EvidenceBank(BaseModel):
    evidence_items: List[EvidenceItem]
    safety_notes: List[str]


class DraftOutput(BaseModel):
    resume_md: str
    linkedin_md: str
    used_evidence_ids: List[str]


class QualityReport(BaseModel):
    passed: bool
    issues: List[str]
    warnings: List[str] = Field(default_factory=list)
    final_resume_md: str
    final_linkedin_md: str


FORBIDDEN_DASHES = ["—", "–"]


def _repo_root() -> Path:
    # /repo/app/pipeline.py -> /repo
    return Path(__file__).resolve().parents[1]


def _read_text(path: Path) -> str:
    if not path.exists():
        return ""
    return path.read_text(encoding="utf-8")


def _truncate_text(text: str, max_chars: int) -> str:
    if not text:
        return ""
    return text[:max_chars]


def _read_docx_text(path: Path, *, max_chars: int = 6000) -> str:
    if not path.exists() or not path.is_file():
        return ""
    doc = Document(str(path))
    parts: List[str] = []
    for para in doc.paragraphs:
        t = (para.text or "").strip()
        if t:
            parts.append(t)
    return "\n".join(parts)[:max_chars]


def _resolve_variant_docx_text(*, variant: str) -> str:
    """
    Accepts variant keys like 'CPTO' and maps them to resumes/*.docx.
    Also accepts a docx relative path (e.g., 'resumes/cpto.docx').
    """
    root = _repo_root()
    resumes_dir = root / "resumes"

    v = (variant or "").strip()
    if not v:
        return ""

    # If it's a path (relative or absolute), try loading directly.
    if "/" in v or v.lower().endswith(".docx"):
        # Treat '/resumes/cpto.docx' and 'resumes/cpto.docx' as repo-relative.
        rel = v.lstrip("/")
        candidate = root / rel
        if candidate.exists():
            return _read_docx_text(candidate)

        # Case-insensitive fallback in resumes/.
        target_name = Path(v).name.lower()
        if resumes_dir.exists():
            for p in resumes_dir.glob("*.docx"):
                if p.name.lower() == target_name:
                    return _read_docx_text(p)
        return ""

    # Otherwise treat as a key.
    key = v.upper()
    key_to_filename = {
        "CPTO": "CPTO.docx",
        "CTO": "CTO.docx",
        "CPO": "CPO.docx",
    }
    desired = key_to_filename.get(key)
    if not desired:
        return ""

    # Case-insensitive match.
    if resumes_dir.exists():
        for p in resumes_dir.glob("*.docx"):
            if p.name.lower() == desired.lower():
                return _read_docx_text(p)

    return ""


def _contains_forbidden_dashes(s: str) -> List[str]:
    hits: List[str] = []
    for dash in FORBIDDEN_DASHES:
        if dash in s:
            hits.append(dash)
    return hits


def _requirements_prompt(*, variant: str, job_description: str, variant_strategy_doc: str) -> str:
    return f"""
You are an Intake and Requirements Agent for an executive resume and LinkedIn profile.

Variant to target: {variant}

Job description (may include multiple roles and mixed formatting):
{job_description}

Variant strategy doc (may contain your preferred positioning, metrics, and leadership themes):
---
{variant_strategy_doc}
---

Output requirements only. Do not write any resume or LinkedIn prose.

Style constraints for downstream agents (must be preserved):
- Hybrid-bullet approach (2 to 3 sentences context, then 2 to 3 metric bullets).
- Executive candor, factual tone, avoid overselling.
- Zero-dash policy: do not use em-dash (—) or en-dash (–) anywhere in narratives. Use '-' only for date ranges and for 'Title - Company' or 'Title | Company' separators.
- ATS realism: extract the keywords to mirror, including leadership and domain terms.
""".strip()


def _evidence_prompt(
    *,
    strategy_v2: str,
    goldmaster: str,
    elevations: str,
    master_resume: str,
    variant_strategy_doc: str,
) -> str:
    return f"""
You are an Evidence Extractor Agent.

Task:
1) Extract executive-grade accomplishment bullets and quantified metrics from the provided documents.
2) Only use numbers and claims that appear explicitly in the text.
3) Return a set of bullet candidates suitable to reuse in a tailored executive resume and LinkedIn profile.

Documents (source of truth):
---
v2_ExecResume_Strategy.md
{strategy_v2}
---
goldmaster_resumes.md
{goldmaster}
---
ElevationsCU.md (board/candidacy narrative and proof points)
{elevations}
---
resumes/MasterResume.md (structure and any non-placeholder content)
{master_resume}
---
Variant strategy doc (may include additional accomplishment language and positioning):
---
{variant_strategy_doc}
---

Output evidence only; do not write full resume sections yet.

Zero-dash policy:
- Never output em-dash (—) or en-dash (–).
- Use '-' only when needed for date ranges and separators between title and company.
""".strip()


def _draft_prompt(*, variant: str, job_description: str, requirements: RequirementsMap, evidence: EvidenceBank) -> str:
    return f"""
You are an Executive Resume Draft Agent and a LinkedIn Profile Draft Agent.

Variant to target: {variant}

Job description:
{job_description}

Requirements map (JSON):
{requirements.model_dump_json()}

Evidence bank (JSON):
{evidence.model_dump_json()}

Write:
1) Executive resume in markdown (ATS-friendly), using Hybrid-bullet approach.
2) LinkedIn profile draft, matching the same evidence and requirements.

Executive formatting rules:
- Use scannable section headers with markdown headings.
- Keep narrative context to 2 to 3 sentences maximum per section where applicable.
- After each context line, provide 2 to 3 bullet lines with quantified impact when available.
- Use bold for key metrics and outcomes (e.g., **66%**).
- Preserve Zero-dash policy: never use em-dash (—) or en-dash (–) anywhere in output.
- Avoid placeholders like '[Paste your current bio here]'; replace with executive-ready placeholders derived from evidence if needed.
- Do not invent metrics not present in the Evidence bank.

Return only structured JSON matching the requested schema.
""".strip()


def _audit_prompt(*, resume_md: str, linkedin_md: str, job_description: str, requirements: RequirementsMap) -> str:
    return f"""
You are an Audit Agent for executive resume and LinkedIn drafts.

Audit goals:
- Enforce Zero-dash policy (no em-dash — or en-dash –).
- Identify missing or weak keyword coverage for ATS and role fit.
- Identify bullets lacking any numeric metric.
- Ensure scannability and formatting consistency.

Job description (for keyword matching):
{job_description}

Requirements map (JSON):
{requirements.model_dump_json()}

Resume draft markdown:
{resume_md}

LinkedIn draft markdown:
{linkedin_md}

Return:
- a list of issues (if any)
- a boolean passed
- revised final_resume_md and final_linkedin_md, fixing issues where possible.

Zero-dash policy:
- Remove any em-dash (—) or en-dash (–).
""".strip()


def _dash_fix_prompt(*, resume_md: str, linkedin_md: str) -> str:
    return f"""
Dash-only enforcement.

Rewrite the provided resume and LinkedIn markdown with these constraints:
- Remove any em-dash (—) and en-dash (–) characters.
- Preserve all other wording, structure, and content as much as possible.
- Follow Zero-dash policy: avoid using '-' as a punctuation substitute; prefer commas or periods instead.

Resume markdown:
{resume_md}

LinkedIn markdown:
{linkedin_md}
""".strip()


def generate_resume_and_linkedin(
    *,
    variant: str,
    job_description_text: str,
    max_input_chars: int = 12000,
) -> Dict[str, Any]:
    jd = _truncate_text(job_description_text or "", max_input_chars)

    root = _repo_root()
    strategy_v2 = _read_text(root / "v2_ExecResume_Strategy.md")
    goldmaster = _read_text(root / "goldmaster_resumes.md")
    elevations = _read_text(root / "ElevationsCU.md")
    master_resume = _read_text(root / "resumes" / "MasterResume.md")
    variant_strategy_doc = _resolve_variant_docx_text(variant=variant)

    model = os.environ.get("GEMINI_MODEL", "gemini-1.5-flash")

    requirements = generate_json(
        schema=RequirementsMap,
        prompt=_requirements_prompt(variant=variant, job_description=jd, variant_strategy_doc=variant_strategy_doc),
        model=model,
    )

    evidence = generate_json(
        schema=EvidenceBank,
        prompt=_evidence_prompt(
            strategy_v2=strategy_v2,
            goldmaster=goldmaster,
            elevations=elevations,
            master_resume=master_resume,
            variant_strategy_doc=variant_strategy_doc,
        ),
        model=model,
    )

    draft = generate_json(
        schema=DraftOutput,
        prompt=_draft_prompt(
            variant=variant,
            job_description=jd,
            requirements=requirements,
            evidence=evidence,
        ),
        model=model,
    )

    # Audit step for final enforcement and revisions.
    audit = generate_json(
        schema=QualityReport,
        prompt=_audit_prompt(
            resume_md=draft.resume_md,
            linkedin_md=draft.linkedin_md,
            job_description=jd,
            requirements=requirements,
        ),
        model=model,
    )

    # Extra local guardrail: catch forbidden dash characters even if the audit misses.
    combined = (audit.final_resume_md or "") + "\n" + (audit.final_linkedin_md or "")
    forbidden_hits = _contains_forbidden_dashes(combined)
    if forbidden_hits:
        dash_fix = generate_json(
            schema=QualityReport,
            prompt=_dash_fix_prompt(
                resume_md=audit.final_resume_md,
                linkedin_md=audit.final_linkedin_md,
            ),
            model=model,
        )
        audit = dash_fix
        audit.issues.append(f"Forbidden dash characters detected in earlier draft: {forbidden_hits}")

    return {
        "resume_md": audit.final_resume_md,
        "linkedin_md": audit.final_linkedin_md,
        "quality_report": {
            "passed": audit.passed,
            "issues": audit.issues,
            "warnings": audit.warnings,
        },
        "audit_raw": audit.model_dump(),
        "used_evidence_ids": draft.used_evidence_ids,
    }

